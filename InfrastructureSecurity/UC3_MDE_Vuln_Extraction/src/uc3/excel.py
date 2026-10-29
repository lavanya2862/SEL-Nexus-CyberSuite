"""Vulnerability Report workbook for one scan (opens in Excel).

data/reporting/Vulnerability_Report_<snapshot>.xlsx. Same numbers as the Word/PDF report (report.gather):
Summary (number tiles + charts), Vulnerabilities (every one, with its remediation beside it), Remediation plan (AI), New this scan,
Remediated this scan, Computers, Top CVEs, Apps to update, Deadlines, Trend, Data quality, How ratings work.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from uc3 import docstyle as ds
from uc3.config import Settings
from uc3.report import AGES, EXPLOIT_WORDS, SEVS, URGENCY_COLOURS, gather, run_lines, summary_tiles

WHITE = "FFFFFF"
STATUS = {"PASS": ds.GOOD, "WARN": ds.WARN, "FAIL": ds.BAD, "INFO": ds.BLUE}
FIX_SHEET = "Remediation plan (AI)"


def _fill(colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=colour)


def _text_on(colour: str) -> str:
    return ds.INK if colour in ds.DARK_TEXT_ON else WHITE


def _band(ws, title: str, subtitle: str, ncols: int) -> None:
    """Navy title band across the top of a sheet, with a blue edge."""
    last = get_column_letter(max(ncols, 6))
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    for row, (text, size, colour) in enumerate(((title, 16, WHITE), (subtitle, 9, "C7D2FE")), 1):
        c = ws.cell(row=row, column=1, value=ds.clean(text))
        c.font = Font(bold=row == 1, size=size, color=colour)
        c.alignment = Alignment(vertical="center", indent=1)
        for col in range(1, max(ncols, 6) + 1):
            ws.cell(row=row, column=col).fill = _fill(ds.NAVY)
    for col in range(1, max(ncols, 6) + 1):
        ws.cell(row=1, column=col).border = Border(top=Side("thick", color=ds.BLUE))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 18
    ws.sheet_view.showGridLines = False


def _table(ws, top: int, name: str, headers: list[str], rows: list[list], widths: list[int],
           colours: dict[int, dict] | None = None, wrap: bool = False, freeze: bool = True) -> int:
    """An Excel table (filters, banded rows) with a navy header; colours = {column index: {value: fill}}.
    Returns the next free row."""
    numeric = {i for i in range(len(headers))
               if isinstance(next((row[i] for row in rows if row[i] not in (None, "")), None), (int, float))}
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=top, column=i, value=h)
        c.font, c.fill = Font(bold=True, color=WHITE), _fill(ds.NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center" if i - 1 in numeric else None)
    ws.row_dimensions[top].height = 30
    for r, row in enumerate(rows, top + 1):
        for i, value in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=ds.clean(value) if isinstance(value, str) else value)
            c.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal="center" if i - 1 in numeric else None)
            if isinstance(value, (int, float)) and not isinstance(value, bool) and abs(value) >= 1000:
                c.number_format = "#,##0"
            fill = (colours or {}).get(i - 1, {}).get(value) if colours else None
            if fill:
                c.fill, c.font = _fill(fill), Font(bold=True, color=_text_on(fill))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(ws.column_dimensions[get_column_letter(i)].width or 0, w)
    if rows:
        ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"
        t = Table(displayName=name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight15", showRowStripes=True)
        ws.add_table(t)
    else:
        ws.cell(row=top + 1, column=1, value="Nothing to show for this scan.").font = Font(italic=True, color=ds.INK_3)
    if freeze:
        ws.freeze_panes = ws.cell(row=top + 1, column=1)
    return top + max(len(rows), 1) + 2


def _tiles(ws, row: int, items: list[tuple], span: int = 2) -> int:
    """Number tiles: value / label / hint in a light box with a coloured top edge, `span` columns wide each."""
    for n, (value, label, hint, colour) in enumerate(items):
        first = 1 + n * span
        last = first + span - 1
        for r, (text, font) in enumerate(((value, Font(bold=True, size=20, color=ds.INK)),
                                          (label, Font(bold=True, size=10, color=ds.INK)),
                                          (hint, Font(size=8.5, color=ds.INK_3)))):
            ws.merge_cells(start_row=row + r, start_column=first, end_row=row + r, end_column=last)
            c = ws.cell(row=row + r, column=first, value=ds.clean(text))
            c.font, c.alignment = font, Alignment(vertical="center", indent=1, wrap_text=r == 2)
            for col in range(first, last + 1):
                cell = ws.cell(row=row + r, column=col)
                cell.fill = _fill(ds.WELL)
                if r == 0:
                    cell.border = Border(top=Side("thick", color=colour or ds.BLUE))
    ws.row_dimensions[row].height, ws.row_dimensions[row + 1].height, ws.row_dimensions[row + 2].height = 34, 18, 26
    return row + 4


def _colour_series(chart, colours: list[str]) -> None:
    for s, colour in zip(chart.series, colours):
        s.graphicalProperties.solidFill = colour
        s.graphicalProperties.line.solidFill = colour


def _doughnut(ws, title: str, labels_ref: Reference, data_ref: Reference, colours: list[str], anchor: str) -> None:
    ch = DoughnutChart(holeSize=55)
    ch.title, ch.height, ch.width = title, 7, 10
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(labels_ref)
    for i, colour in enumerate(colours):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = colour
        pt.graphicalProperties.line.solidFill = WHITE
        ch.series[0].dPt.append(pt)
    ws.add_chart(ch, anchor)


def _bars(ws, title: str, data_ref: Reference, cats_ref: Reference, colours: list[str], anchor: str,
          horizontal=False, stacked=True, width=16, height=7.5) -> None:
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title, ch.width, ch.height = title, width, height
    if stacked:
        ch.grouping, ch.overlap = "stacked", 100
    ch.add_data(data_ref, titles_from_data=True)
    ch.set_categories(cats_ref)
    _colour_series(ch, colours)
    ws.add_chart(ch, anchor)


def _steps(items) -> str:
    return "\n".join(f"{n}. {ds.clean(s)}" for n, s in enumerate(items or [], 1))


def _yes(flag) -> str:
    return "Yes" if flag else ""


def build_excel(settings: Settings, d: dict | None = None) -> Path:
    d = d or gather(settings)
    t, fx = d["totals"], d["fixes"]
    subtitle = "  ·  ".join(run_lines(d)[:2])
    wb = Workbook()
    sev_colours = {s: ds.SEVERITY[s] for s in (*SEVS, "Unrated")}

    # ---- Fix plan rows first: the Vulnerabilities sheet links each row to its guide
    apps = sorted(fx.get("apps", []), key=lambda a: (a["urgency"]["level"], -a["findings"]))
    settings_cards = sorted(fx.get("settings", []), key=lambda s: ((s.get("urgency") or {}).get("level", 9), -s.get("impact", 0)))
    guide_row = {a["key"]: 5 + i for i, a in enumerate(apps)}

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    _band(ws, "Vulnerability Report", subtitle, 12)
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 14
    row = _tiles(ws, 4, summary_tiles(t))
    row = _tiles(ws, row, [(f"{t['overdue']:,}", "Overdue", "past their due date", ds.BAD),
                           (f"{t['p1']:,}", "Priority P1", "remediate first", ds.SERIOUS),
                           (f"{t['app_versions']:,}", "App versions to update", "each clears several findings", ds.BLUE),
                           (f"{len(apps) + len(settings_cards):,}", "Remediation guides", "one per app or setting", ds.BLUE)])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
    c = ws.cell(row=row, column=1, value=(
        f"This scan found {t['unique']:,} vulnerabilities on {t['machines']} computers. "
        f"{t['critical']:,} are Critical and {t['high']:,} High; "
        f"{t['overdue']:,} are past their fix-by date. Since the last scan, {t['new']:,} appeared and {t['fixed']:,} were fixed. "
        f"Every vulnerability is listed on the 'Vulnerabilities' sheet with its remediation."))
    c.alignment, c.font = Alignment(wrap_text=True, vertical="top"), Font(size=10, color=ds.INK)
    ws.row_dimensions[row].height = 44
    row += 2

    # chart data blocks (left) with their charts (right)
    sev_names = [*SEVS, "Unrated"]
    top = row
    ws.cell(row=top, column=1, value="Severity").font = Font(bold=True, color=WHITE)
    ws.cell(row=top, column=2, value="Vulnerabilities").font = Font(bold=True, color=WHITE)
    for col in (1, 2):
        ws.cell(row=top, column=col).fill = _fill(ds.NAVY)
    for i, s in enumerate(sev_names, 1):
        ws.cell(row=top + i, column=1, value=s).fill = _fill(ds.SEVERITY[s])
        ws.cell(row=top + i, column=1).font = Font(bold=True, color=_text_on(ds.SEVERITY[s]))
        ws.cell(row=top + i, column=2, value=d["sev"].get(s, 0)).number_format = "#,##0"
    _doughnut(ws, "By severity", Reference(ws, min_col=1, min_row=top + 1, max_row=top + len(sev_names)),
              Reference(ws, min_col=2, min_row=top, max_row=top + len(sev_names)),
              [ds.SEVERITY[s] for s in sev_names], f"D{top}")
    prio = ("P1", "P2", "P3", "P4")
    ptop = top + len(sev_names) + 2
    ws.cell(row=ptop, column=1, value="Priority").font = Font(bold=True, color=WHITE)
    ws.cell(row=ptop, column=2, value="Vulnerabilities").font = Font(bold=True, color=WHITE)
    for col in (1, 2):
        ws.cell(row=ptop, column=col).fill = _fill(ds.NAVY)
    for i, p in enumerate(prio, 1):
        ws.cell(row=ptop + i, column=1, value=p)
        ws.cell(row=ptop + i, column=2, value=d["prio"].get(p, 0)).number_format = "#,##0"
    _doughnut(ws, "By priority", Reference(ws, min_col=1, min_row=ptop + 1, max_row=ptop + 4),
              Reference(ws, min_col=2, min_row=ptop, max_row=ptop + 4), [ds.BAD, ds.SERIOUS, ds.WARN, "16A34A"], f"I{top}")
    otop = ptop + 7
    heads = ["OpCo", *SEVS, "Total", "P1", "Overdue"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(row=otop, column=i, value=h)
        cell.font, cell.fill = Font(bold=True, color=WHITE), _fill(ds.NAVY)
    for r, o in enumerate(d["by_opco"], otop + 1):
        for i, v in enumerate([o["opco"], o["critical"], o["high"], o["medium"], o["low"], o["total"], o["p1"], o["overdue"]], 1):
            ws.cell(row=r, column=i, value=v)
    if d["by_opco"]:
        _bars(ws, "Vulnerabilities by OpCo and severity",
              Reference(ws, min_col=2, min_row=otop, max_col=5, max_row=otop + len(d["by_opco"])),
              Reference(ws, min_col=1, min_row=otop + 1, max_row=otop + len(d["by_opco"])),
              [ds.SEVERITY[s] for s in SEVS], f"A{otop + len(d['by_opco']) + 2}", width=24)

    # ---- Vulnerabilities: every one, with its fix
    ws = wb.create_sheet("Vulnerabilities")
    headers = ["Computer", "OpCo", "App", "Version", "CVE", "Severity", "CVSS", "Priority", "Exploit",
               "First seen", "Open (days)", "Remediation SLA (days)", "Due date", "Overdue",
               "Recommended update", "Update available", "Priority", "Remediation", "Command", "What is wrong", "Steps",
               "How to check", "Remediation guide"]
    _band(ws, f"All vulnerabilities in this scan ({len(d['vulns']):,})", subtitle + "  ·  each row carries its remediation; "
          "'Remediation guide' opens the full guide", len(headers))
    rows = []
    for v in d["vulns"]:
        g = v["guide"] or {}
        rows.append([v["device_name"], v["opco"], f"{v['software_vendor']} {v['software_name']}", v["software_version"],
                     v["cve_id"], v["severity"], v["cvss_score"], v["priority"], v["exploit_words"],
                     v["first_seen"], v["age_days"], v["sla_days"], v["fix_by"], _yes(v["sla_breached"]),
                     v["recommended_update"] or "", _yes(v["security_update_available"]),
                     (g.get("urgency") or {}).get("label", ""), (g.get("fix") or {}).get("action", ""),
                     (g.get("fix") or {}).get("command") or "", (g.get("ai") or {}).get("what_is_wrong", ""),
                     _steps((g.get("ai") or {}).get("steps")).replace("\n", "   "), (g.get("ai") or {}).get("verify", ""),
                     g.get("app", "") and "Open guide"])
    _table(ws, 4, "Vulnerabilities", headers, rows,
           [14, 10, 22, 16, 17, 10, 7, 8, 15, 12, 9, 9, 12, 9, 20, 10, 24, 34, 34, 50, 60, 40, 12],
           {5: sev_colours, 13: {"Yes": ds.BAD}})
    for r, v in enumerate(d["vulns"], 5):
        ws.row_dimensions[r].height = 16                 # one line per vulnerability; the full guide is one click away
        for col in (10, 13):
            ws.cell(row=r, column=col).number_format = "dd mmm yyyy"
        g = v["guide"]
        if g and g["key"] in guide_row:
            c = ws.cell(row=r, column=23)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{FIX_SHEET}'!A{guide_row[g['key']]}")
            c.font = Font(color=ds.BLUE, underline="single")
    ws.freeze_panes = "F5"

    # ---- Remediation plan (AI)
    ws = wb.create_sheet(FIX_SHEET)
    headers = ["#", "Kind", "Priority", "App / setting", "Versions / category", "Remediation type", "Due date", "Computers", "CVEs",
               "Critical", "High", "Remediation", "Command", "What is wrong", "Worst case", "Why now", "Steps",
               "How to check", "Side effects", "Written by"]
    _band(ws, f"Remediation plan: {len(apps)} apps and {len(settings_cards)} security settings",
          "Facts, versions, deadlines and commands come from the data; the AI explains them. Check before acting.", len(headers))

    def by(g: dict) -> str:
        return f"AI ({g.get('model') or '?'})" if (g.get("source") or "").startswith("ai") else "Data only"

    rows = [[i, "App", a["urgency"]["label"], a["app"], ", ".join(v["version"] for v in a["versions"]), a.get("fix_type"),
             date.fromisoformat(a["fix_by"]) if a.get("fix_by") else None, ", ".join(x["name"] for x in a["devices"]), a["cves"], a["critical"], a["high"],
             (a.get("fix") or {}).get("action"), (a.get("fix") or {}).get("command") or "",
             a["ai"].get("what_is_wrong"), a["ai"].get("worst_case"), a["ai"].get("why_now"), _steps(a["ai"].get("steps")),
             a["ai"].get("verify"), a["ai"].get("side_effects"), by(a)] for i, a in enumerate(apps, 1)]
    rows += [[len(apps) + i, "Setting", (s.get("urgency") or {}).get("label", ""), s["app"],
              f"{s.get('category', '')} / {s.get('subcategory', '')} · impact {s.get('impact', 0):g}/10", s.get("fix_type"), None,
              ", ".join(x["name"] for x in s.get("devices", [])), None, None, None, (s.get("fix") or {}).get("action"),
              (s.get("fix") or {}).get("command") or "", s["ai"].get("what_is_wrong"), s["ai"].get("worst_case"),
              s["ai"].get("why_now"), _steps(s["ai"].get("steps")), s["ai"].get("verify"), s["ai"].get("side_effects"), by(s)]
             for i, s in enumerate(settings_cards, 1)]
    urgency_fill = {g["urgency"]["label"]: URGENCY_COLOURS.get(g["urgency"]["level"], ds.BLUE)
                    for g in [*apps, *settings_cards] if g.get("urgency")}
    _table(ws, 4, "FixPlan", headers, rows,
           [5, 8, 22, 30, 26, 22, 11, 22, 7, 8, 7, 32, 36, 50, 40, 40, 70, 40, 34, 16], {2: urgency_fill}, wrap=True)
    for r in range(5, 5 + len(apps)):
        ws.cell(row=r, column=7).number_format = "dd mmm yyyy"
    ws.freeze_panes = "E5"

    # ---- New / Fixed this scan
    for status, title, table_name in (("New", "New in this scan", "NewThisScan"), ("Fixed", "Remediated in this scan", "FixedThisScan")):
        ws = wb.create_sheet(title)
        items = d["changes"][status]
        headers = ["Computer", "OpCo", "App", "Version", "CVE", "Severity", "CVSS", "Exploit", "Recommended update",
                   "When (UTC)"]
        _band(ws, f"{title} ({len(items):,})", "From Defender's own change list since the last scan", len(headers))
        _table(ws, 4, table_name, headers,
               [[x["device"], x["opco"], x["app"], x["version"], x["cve"], x["severity"], x["cvss"], x["exploit"],
                 x["update"], x["when"]] for x in items],
               [14, 10, 24, 16, 17, 10, 7, 16, 22, 17], {5: sev_colours})

    # ---- Computers
    ws = wb.create_sheet("Computers")
    per_device: dict[str, set] = {}
    for v in d["vulns"]:
        if v["guide"]:
            per_device.setdefault(v["device_name"], set()).add(v["guide"]["app"])
    headers = ["Computer", "OpCo", "OS", "Findings", *SEVS, "P1", "Overdue", "App versions", "Exposure", "Risk score",
               "Health", "Last seen", "Apps to update"]
    _band(ws, f"Computers ({len(d['devices'])})", "Findings per computer and the apps to update on it", len(headers))
    _table(ws, 4, "Computers", headers,
           [[x["device_name"], x["opco"], f"{x['os_platform'] or ''} {x['os_version'] or ''}".strip(), x["findings"],
             x["critical"], x["high"], x["medium"], x["low"], x["p1"], x["overdue"], x["app_versions"], x["exposure_level"],
             x["risk_score"], x["health_status"], x["last_seen"], ", ".join(sorted(per_device.get(x["device_name"], [])))]
            for x in d["devices"]],
           [16, 10, 22, 10, 9, 9, 9, 9, 8, 9, 11, 11, 11, 11, 12, 80], wrap=True)

    # ---- Top CVEs
    ws = wb.create_sheet("Top CVEs")
    headers = ["CVE", "Severity", "CVSS", "Priority", "Exploit", "Computers", "Findings", "Apps", "OpCos"]
    _band(ws, "Top 100 CVEs", "Ordered by priority, then how many computers have it", len(headers))
    _table(ws, 4, "TopCVEs", headers,
           [[c["cve_id"], c["severity"], c["cvss"], c["priority"],
             EXPLOIT_WORDS.get(c["exploitability"], c["exploitability"] or ""), c["devices"], c["findings"], c["apps"], c["opcos"]]
            for c in d["top_cves"]],
           [17, 10, 7, 9, 16, 10, 9, 40, 18], {1: sev_colours})

    # ---- Apps to update
    ws = wb.create_sheet("Apps to update")
    guides = {a["key"]: a for a in apps}
    headers = ["App", "Version", *SEVS, "Findings cleared", "Computers", "Priority", "Remediation", "Command", "Recommended update"]
    _band(ws, f"App versions to update ({len(d['apps'])})", "Updating one app version clears all of its findings", len(headers))
    rows = []
    for a in d["apps"]:
        g = guides.get(f"{a['software_vendor']}:{a['software_name']}") or {}
        rows.append([f"{a['software_vendor']} {a['software_name']}", a["software_version"], a["critical"], a["high"],
                     a["medium"], a["low"], a["findings"], a["devices"], (g.get("urgency") or {}).get("label", ""),
                     (g.get("fix") or {}).get("action", ""), (g.get("fix") or {}).get("command") or "", a["update"] or ""])
    _table(ws, 4, "AppsToUpdate", headers, rows, [26, 16, 9, 9, 9, 9, 10, 10, 24, 36, 36, 22], {8: urgency_fill})
    n = min(len(d["apps"]), 15)
    if n:
        ws.cell(row=4, column=14, value="Top 15 (chart data)").font = Font(bold=True)
        for i, a in enumerate(d["apps"][:n], 5):
            ws.cell(row=i, column=14, value=f"{a['software_name']} {a['software_version']}")
            ws.cell(row=i, column=15, value=a["findings"])
        ws.cell(row=4, column=15, value="Findings")
        _bars(ws, "Findings cleared by each update (top 15)", Reference(ws, min_col=15, min_row=4, max_row=4 + n),
              Reference(ws, min_col=14, min_row=5, max_row=4 + n), [ds.BLUE], "Q4", horizontal=True, stacked=False,
              width=18, height=10)

    # ---- Deadlines
    ws = wb.create_sheet("Deadlines")
    _band(ws, "Deadlines and age", "Due date = first seen + the deadline for its severity", 10)
    nxt = _table(ws, 4, "Deadlines", ["Severity", "Remediation SLA (days)", "Open", "Overdue"],
                 [[s["severity"], s["days"] if s["days"] is not None else "no deadline", s["open"], s["overdue"]] for s in d["sla"]],
                 [16, 16, 12, 12], {0: sev_colours}, freeze=False)
    age_top = nxt + 1
    _table(ws, age_top, "Ages", ["Open for (days)", *SEVS],
           [[a, *[(d["ages"].get(a) or {}).get(s.lower(), 0) for s in SEVS]] for a in AGES], [16, 16, 12, 12, 12], freeze=False)
    _bars(ws, "How long findings have been open", Reference(ws, min_col=2, min_row=age_top, max_col=5, max_row=age_top + 4),
          Reference(ws, min_col=1, min_row=age_top + 1, max_row=age_top + 4), [ds.SEVERITY[s] for s in SEVS], "G4")
    old_top = age_top + 8
    ws.cell(row=old_top - 1, column=1, value="Oldest Critical and High findings").font = Font(bold=True, size=11, color=ds.NAVY)
    _table(ws, old_top, "Oldest", ["CVE", "App", "Computer", "Severity", "Open (days)", "Due date", "Overdue"],
           [[v["cve_id"], f"{v['software_name']} {v['software_version']}", v["device_name"], v["severity"], v["age_days"],
             v["fix_by"], _yes(v["sla_breached"])] for v in d["oldest"]],
           [16, 16, 12, 12, 12, 12, 10], {3: sev_colours, 6: {"Yes": ds.BAD}}, freeze=False)

    # ---- Trend
    ws = wb.create_sheet("Trend")
    _band(ws, "Open vulnerabilities per scan", "One row per scan day", 8)
    _table(ws, 4, "Trend", ["Scan", *SEVS], [[x["day"], *[x[s.lower()] for s in SEVS]] for x in d["trend"]],
           [12, 10, 10, 10, 10], freeze=False)
    if len(d["trend"]) > 1:
        ch = LineChart()
        ch.title, ch.width, ch.height = "Open vulnerabilities per scan", 22, 9
        ch.add_data(Reference(ws, min_col=2, min_row=4, max_col=5, max_row=4 + len(d["trend"])), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(d["trend"])))
        for s, colour in zip(ch.series, [ds.SEVERITY[x] for x in SEVS]):
            s.graphicalProperties.line.solidFill = colour
            s.graphicalProperties.line.width = 22000
            s.smooth = False
        ws.add_chart(ch, "G4")

    # ---- Data quality
    ws = wb.create_sheet("Data quality")
    _band(ws, f"Data quality: {t['quality']}", "Automatic checks on this scan, then every scan so far", 8)
    nxt = _table(ws, 4, "Checks", ["Check", "Result", "Value", "Expected", "What it means"],
                 [[c["name"].replace("_", " "), c["status"], c["value"], c["expected"], c.get("message", "")]
                  for c in d["quality"].get("checks", [])], [24, 10, 28, 34, 70], {1: STATUS}, wrap=True, freeze=False)
    ws.cell(row=nxt, column=1, value="Scan history").font = Font(bold=True, size=11, color=ds.NAVY)
    _table(ws, nxt + 1, "History", ["Scan", "Rows in Defender", "Rows downloaded", "Repeats", "Vulnerabilities kept",
                                    "Coverage", "Computers", "Run time (s)", "Quality"],
           [[h["run_date"], h["defender_count"], h["vulnerability_rows"], h.get("repeats"), h.get("vulnerabilities"),
             f"{h['coverage']:.1%}" if h["coverage"] is not None else "", h["machines"], h["duration_seconds"], h["quality"]]
            for h in d["history"]], [24, 10, 28, 34, 12, 12, 12, 12, 10], {8: STATUS}, freeze=False)

    # ---- How ratings work
    ws = wb.create_sheet("How ratings work")
    _band(ws, "How the ratings work", "Read this once; every sheet uses these rules", 6)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110
    for r, (k, v) in enumerate([
        ("Severity", "Comes from Microsoft Defender, mostly from the CVE's CVSS score. A brand-new CVE Microsoft has not rated "
                     "yet is Unrated."),
        ("Priority P1", "Critical, or High with a verified exploit or one in attack kits."),
        ("Priority P2", "High, or Medium with such an exploit."),
        ("Priority P3 / P4", "P3 = Medium or Unrated; P4 = Low."),
        ("Due date", "The day Defender first saw it + the deadline for its severity: Critical 15 days, High 30, Medium 90, "
                   "Low 180. After that it is overdue."),
        ("New / Remediated", "From Defender's own change list for this scan."),
        ("Repeats", "Rows Defender sent twice in its export; counted once."),
        ("Remediation guides", "One per app and per weak setting. Versions, CVEs, computers and commands come from the data; the AI "
                       "only explains them."),
    ], 4):
        ws.cell(row=r, column=1, value=k).font = Font(bold=True, color=ds.NAVY)
        ws.cell(row=r, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30

    out = settings.reporting_dir / f"Vulnerability_Report_{d['snapshot']}.xlsx"
    wb.save(out)
    return out
