"""AI spot-check: a person checks a random sample of Gemini's answers against the source, which gives an
honest 'AI accuracy' figure for the demo.

  python -m threat_intel spotcheck                    -> data/reports/Threat_Intel_AI_Spotcheck_<date>.xlsx (20 random AI answers)
  python -m threat_intel spotcheck --score <file>     -> reads your Y/N marks, writes data/ai_spotcheck_result.json
"""

from __future__ import annotations

import json
import random
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from threat_intel.config import Settings
from threat_intel.store import Store

CHECK_COLUMNS = ["Summary correct? (Y/N)", "Products correct? (Y/N)", "CVEs correct? (Y/N)", "Actions sensible? (Y/N)"]
HEADERS = ["#", "Threat ID", "Source", "Title", "Link to the source", "Source text (start)", "AI summary",
           "AI products", "AI CVEs", "AI techniques", "AI actions", *CHECK_COLUMNS, "Notes"]
INSTRUCTIONS = [
    "What this is: 20 random threats the AI (Gemini) analysed. You check whether the AI got the facts right.",
    "How: for each row, read the 'Source text' (or open the link) and compare it with the AI columns.",
    "Mark each check column Y (right) or N (wrong). Leave it empty if you can't tell.",
    "Summary correct? = the summary says nothing that isn't in the source and misses nothing important.",
    "Products correct? = the affected products/vendors match the source (empty is fine if the source names none).",
    "CVEs correct? = the CVE IDs match the source (empty is fine if the source names none).",
    "Actions sensible? = the recommended actions follow from the source (e.g. the right version to update to).",
    "You don't need security knowledge: it's 'does the AI's text match the source text?'.",
    "When done: save the file and run  .venv\\Scripts\\python.exe -m threat_intel spotcheck --score \"<this file>\"",
]


def _fmt(value) -> str:
    items = json.loads(value) if isinstance(value, str) else (value or [])
    out = []
    for i in items:
        if isinstance(i, dict):
            out.append(" ".join(str(v) for v in i.values() if v))
        else:
            out.append(str(i))
    return "\n".join(out)


def export(store: Store, settings: Settings, n: int = 20, seed: int | None = None) -> Path:
    rows = store.query("""
        SELECT t.threat_id, t.source, t.title, t.url, t.text, a.summary, a.products, a.techniques, a.actions,
               (SELECT string_agg(value, ', ') FROM threat_links l WHERE l.threat_id = t.threat_id
                  AND l.link_type = 'cve' AND l.origin = 'ai') AS ai_cves
        FROM threats t JOIN analysis a USING (threat_id) WHERE a.ai_status = 'done' AND NOT t.is_test""")
    random.Random(seed).shuffle(rows)
    sample = rows[:n]
    wb = Workbook()
    ws = wb.active
    ws.title = "Check these"
    ws.append(HEADERS)
    for i, r in enumerate(sample, start=1):
        ws.append([i, r["threat_id"], r["source"], r["title"], r["url"] or "", (r["text"] or "")[:1500], r["summary"],
                   _fmt(r["products"]), r["ai_cves"] or "", _fmt(r["techniques"]), _fmt(r["actions"]),
                   "", "", "", "", ""])
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [4, 22, 12, 40, 30, 60, 60, 30, 18, 30, 50, 14, 14, 14, 14, 30]
    for col, width in zip("ABCDEFGHIJKLMNOP", widths):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(yn)
    yn.add(f"L2:O{len(sample) + 1}")
    ws.freeze_panes = "B2"
    info = wb.create_sheet("How to check")
    for line in INSTRUCTIONS:
        info.append([line])
    info.column_dimensions["A"].width = 120
    settings.reports_dir.mkdir(parents=True, exist_ok=True)
    path = settings.reports_dir / f"Threat_Intel_AI_Spotcheck_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return path


def score(path: Path, settings: Settings) -> dict:
    ws = load_workbook(path, data_only=True)["Check these"]
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in CHECK_COLUMNS}
    per_check = {name: {"Y": 0, "N": 0} for name in CHECK_COLUMNS}
    for row in ws.iter_rows(min_row=2, values_only=True):
        for name, i in idx.items():
            mark = str(row[i] or "").strip().upper()[:1]
            if mark in ("Y", "N"):
                per_check[name][mark] += 1
    total_y = sum(v["Y"] for v in per_check.values())
    total = sum(v["Y"] + v["N"] for v in per_check.values())
    result = {
        "file": str(path), "scored_at": datetime.now().isoformat(timespec="seconds"),
        "checks": {k: {**v, "pct_correct": round(100 * v["Y"] / (v["Y"] + v["N"]), 1) if v["Y"] + v["N"] else None}
                   for k, v in per_check.items()},
        "overall_pct_correct": round(100 * total_y / total, 1) if total else None, "marks": total,
    }
    (settings.data_dir / "ai_spotcheck_result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result


def latest_result(settings: Settings) -> dict | None:
    path = settings.data_dir / "ai_spotcheck_result.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None
